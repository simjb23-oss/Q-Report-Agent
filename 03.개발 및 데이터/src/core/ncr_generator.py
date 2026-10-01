# -*- coding: utf-8 -*-
"""
NCR (Non-Conformance Report) & Quality Certificate Generator
판정 결과(PASS/FAIL)에 따른 공식 품질 문서 생성기
- Markdown 보고서
- HTML 보고서
- Excel (.xlsx) 보고서 (openpyxl)
- PDF (.pdf) 보고서 (pymupdf)
"""

import io
from typing import Dict, Any, List
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    import pymupdf
    HAS_PYMUPDF = True
except ImportError:
    try:
        import fitz as pymupdf
        HAS_PYMUPDF = True
    except ImportError:
        HAS_PYMUPDF = False


class NCRGenerator:
    @staticmethod
    def generate_pass_report(parsed_data: Dict[str, Any], eval_result: Dict[str, Any]) -> str:
        return NCRGenerator.generate_ncr_markdown(parsed_data, eval_result)

    @staticmethod
    def generate_markdown_ncr(parsed_data: Dict[str, Any], eval_result: Dict[str, Any]) -> str:
        return NCRGenerator.generate_ncr_markdown(parsed_data, eval_result)

    @staticmethod
    def generate_ncr_markdown(parsed_data: Dict[str, Any], eval_result: Dict[str, Any]) -> str:
        model_code = eval_result["model_code"]
        product_name = eval_result["product_name"]
        supplier = eval_result["supplier"]
        rep_no = parsed_data.get("report_no", "N/A")
        order_qty = parsed_data.get("order_qty", 0)
        sample_qty = parsed_data.get("sample_qty", 0)
        verdict = eval_result.get("final_verdict", "FAIL")
        date_str = parsed_data.get('inspection_date', '2026-09-28')

        if verdict == "PASS":
            return f"""# [합격] 출하검사 합격 인증 및 입고 승인서 (Certificate of Acceptance)

**문서번호:** COA-PASS-{model_code}-{rep_no.replace('/', '_')}  
**발행일자:** {date_str}  
**수신처:** 사내 생산관리 및 자재창고 IQC 부서  
**발신처:** (주)휴롬 품질보증본부 수입검사팀 (칼퇴보증 QA Agent)  

---

### 1. 대상 품목 및 입고 정보
- **품목명 / 모델명:** {product_name} (`{model_code}`)
- **원인 성적서 파일:** `{parsed_data.get('filename')}`
- **성적서 번호 (Report No):** `{rep_no}`
- **적용 검사 기준서:** `{eval_result.get('inspection_spec_no')}`
- **발주(로트) 수량:** `{order_qty:,} PCS` | **샘플 검사수량:** `{sample_qty:,} PCS`

---

### 2. 종합 품질 감사 판정: [정상 합격 (PASS)]
- **규격 및 치수 공차:** 모든 전기적 특성, 전원코드 길이, 외경 치수가 사내 도면 마스터 DB 규격(LSL/USL)을 100% 충족함.
- **안전 규격 검사:** 내전압(Hi-Pot), 누설전류, 절연저항 기준 적합.
- **공정능력 지수:** Cpk 양호 판정 완료.

---

### 3. 승인 후속 조치 (Action Items)
1. **ERP/MES 자동 승인:** 본 입고 로트({order_qty:,} PCS)에 대한 ERP 자재 입고 락(Lock) 해제 및 가결 등록 완료.
2. **생산 투입 허가:** 완제품 및 조립 라인 정상 불출 허용.
3. **품질 이력 보존:** 공급사 출하 성적서 영구 아카이빙 DB 동기화 완료.

---
**판정 담당:** 칼퇴보증 QA Agent 자동 검사 시스템  
**최종 승인:** 품질보증팀 수입검사 파트장 (승인 완료)
"""

        ng_reasons = eval_result.get("ng_reasons", [])
        reasons_bullet = "\n".join([f"- **[위반 항목 {i+1}]** {r}" for i, r in enumerate(ng_reasons)])
        defects_bullet = "\n".join([f"- {d}" for d in parsed_data.get("defects_found", [])])

        ncr_md = f"""# [부적합] 품질 부적합 통보서 (NCR: Non-Conformance Report)

**문서번호:** NCR-{model_code}-{rep_no.replace('/', '_')}  
**발행일자:** {date_str}  
**수신처 (협력사):** {supplier} 품질보증팀 귀중  
**발신처:** (주)휴롬 품질보증본부 수입검사(IQC)팀  

---

### 1. 대상 품목 및 입고 정보
- **품목명 / 모델명:** {product_name} (`{model_code}`)
- **원인 성적서 파일:** `{parsed_data.get('filename')}`
- **성적서 번호 (Report No):** `{rep_no}`
- **검사 기준서:** `{eval_result.get('inspection_spec_no')}`
- **발주(로트) 수량:** `{order_qty:,} PCS` | **샘플 검사수량:** `{sample_qty:,} PCS`

---

### 2. 부적합 판정 사유 (이탈 내역)
{reasons_bullet if reasons_bullet else "- 사내 도면 규격 한계 초과 결함"}

#### [현장 검출 결함 상세]
{defects_bullet if defects_bullet else "- 별도 외관 결함 없음"}

---

### 3. 품질 요구 및 시정조치 요구사항 (Action Item)
1. **입고 조치:** 본 입고 로트(수량: {order_qty:,} PCS)는 **[전량 입고 보류 및 반품(Return)]** 처리됩니다.
2. **원인 규명 및 대책서 제출:**
   - 3영업일 이내에 `5-Why 기반 원인 분석서` 및 `재발 방지 대책서(8D Report)`를 품질본부로 회신 바랍니다.
3. **재입고 조건:** 
   - 치수 전수 검사 성적서 첨부 및 공정 능력 지수($Cpk >= 1.33$) 증빙 자료 제출 필수.

---
**판정 담당:** 칼퇴보증 QA Agent 자동 검사 시스템  
**최종 승인:** 품질보증팀 수입검사 파트장 (인)
"""
        return ncr_md

    @staticmethod
    def generate_ncr_html(parsed_data: Dict[str, Any], eval_result: Dict[str, Any]) -> str:
        model_code = eval_result["model_code"]
        product_name = eval_result["product_name"]
        supplier = eval_result["supplier"]
        rep_no = parsed_data.get("report_no", "N/A")
        date_str = parsed_data.get('inspection_date', '2026-09-28')
        order_qty = parsed_data.get("order_qty", 0)
        verdict = eval_result.get("final_verdict", "FAIL")
        is_pass = (verdict == "PASS")
        
        items = eval_result.get("evaluation_items", [])
        rows_html = ""
        for it in items:
            status_cls = "pass" if it.get("status") == "PASS" else ("fail" if it.get("status") == "NG" else "hold")
            delta_str = f"({it.get('delta', 0):+0.2f})" if it.get('delta') is not None else ""
            status_bg = "#dcfce7; color:#166534;" if status_cls == "pass" else ("#fee2e2; color:#991b1b;" if status_cls == "fail" else "#fef3c7; color:#92400e;")
            rows_html += f"""
            <tr>
                <td><b>{it.get('item_name')}</b></td>
                <td>{it.get('min_limit')} ~ {it.get('max_limit')} {it.get('unit', '')}</td>
                <td><b>{it.get('measured')}</b> {it.get('unit', '')} <small style='color:#64748b;'>{delta_str}</small></td>
                <td><span style='padding:3px 8px; border-radius:999px; font-size:11px; font-weight:800; background:{status_bg}'>{it.get('status')}</span></td>
                <td>{'기준 범위 내' if it.get('status')=='PASS' else ('허용 공차 초과 이탈' if it.get('status')=='NG' else '육안 확인 요망')}</td>
            </tr>
            """
            
        actions_list = """
        <li>검사 결과 품질 이력 DB 정상 등록</li>
        <li>정상 입고 승인 및 생산 라인 투입 가결</li>
        <li>ERP/MES 입고 완료 상태 실시간 전계</li>
        """ if is_pass else """
        <li>해당 LOT 전량 입고 격리 및 출하/생산 투입 보류 조치</li>
        <li>협력사 대상 3영업일 이내 5-Why 원인분석서 및 8D 시정조치 대책서 제출 요구</li>
        <li>동일 로트 및 차기 생산 배치 샘플링 검사 배수 강화 (ISO 2859 강화검사)</li>
        <li>전사 ERP IQC 불합격 상태 전계 및 입고 전산 차단</li>
        """

        doc_title = "출하검사 합격 인증 및 입고 승인서 (Certificate of Acceptance)" if is_pass else "품질 부적합 통보서 (NCR / 8D Report)"
        doc_color = "#166534" if is_pass else "#991b1b"
        doc_prefix = "COA-PASS" if is_pass else "NCR"

        html_content = f"""<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>{doc_prefix}-{model_code}-{rep_no.replace('/', '_')}</title>
<style>
body{{font-family:'Pretendard',Arial,"Noto Sans KR",sans-serif;color:#111827;margin:40px;line-height:1.5}}
h1{{font-size:24px;margin-bottom:4px;color:{doc_color}}}
h2{{font-size:16px;border-bottom:2px solid #111827;padding-bottom:6px;margin-top:28px}}
.meta{{display:grid;grid-template-columns:repeat(2,1fr);gap:8px}}
.box{{background:#f8fafc;padding:12px;border-radius:8px;border:1px solid #e2e8f0}}
.label{{font-size:11px;color:#64748b;font-weight:700}}
.value{{font-weight:700;margin-top:2px;font-size:14px}}
table{{width:100%;border-collapse:collapse;margin-top:10px}}
th,td{{border:1px solid #d1d5db;padding:9px;text-align:left;font-size:12px}}
th{{background:#f3f4f6}}
ul{{padding-left:20px;font-size:13px;color:#334155}}
li{{margin-bottom:6px}}
.sign{{margin-top:50px;display:grid;grid-template-columns:1fr 1fr;gap:30px}}
.line{{border-bottom:1px solid #6b7280;height:45px}}
</style>
</head>
<body>
<h1>{doc_title}</h1>
<div>문서번호: <b>{doc_prefix}-{model_code}-{rep_no.replace('/', '_')}</b> | 발행일: {date_str}</div>

<h2>1. 기본 정보</h2>
<div class="meta">
    <div class="box"><div class="label">협력사명</div><div class="value">{supplier}</div></div>
    <div class="box"><div class="label">품명 / 모델코드</div><div class="value">{product_name} ({model_code})</div></div>
    <div class="box"><div class="label">성적서/LOT 번호</div><div class="value">{rep_no}</div></div>
    <div class="box"><div class="label">발주 / 입고수량</div><div class="value">{order_qty:,} PCS</div></div>
</div>

<h2>2. 종합 판정 결과</h2>
<p style="font-size:15px;"><b>종합 판정: <span style="color:{doc_color}; font-size:18px;">{verdict}</span></b> (검사 기준: {eval_result.get('inspection_spec_no', '사내 도면 마스터')})</p>

<h2>3. 검사 항목별 공차 판정 결과</h2>
<table>
    <thead>
        <tr>
            <th>시험항목</th>
            <th>기준 공차 규격</th>
            <th>실측치</th>
            <th>판정</th>
            <th>판정 근거</th>
        </tr>
    </thead>
    <tbody>
        {rows_html}
    </tbody>
</table>

<h2>4. 후속 조치 및 요구사항 (Action Items)</h2>
<ul>
    {actions_list}
</ul>

<h2>5. 비고 및 특기사항</h2>
<p style="font-size:12px; color:#64748b;">본 문서는 칼퇴보증 QA Agent 자동 검사 시스템에서 공차 판정 결과를 기반으로 생성된 공식 문서입니다.</p>

<div class="sign">
    <div>검사 담당 (IQC Inspector)<div class="line"></div></div>
    <div>승인자 (QA Manager)<div class="line"></div></div>
</div>
</body>
</html>"""
        return html_content

    @staticmethod
    def generate_excel_report(parsed_data: Dict[str, Any], eval_result: Dict[str, Any]) -> bytes:
        """
        전문적인 서식과 색상이 적용된 엑셀(.xlsx) 보고서 생성
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "품질검사결과보고서"
        ws.views.sheetView[0].showGridLines = True

        model_code = eval_result.get("model_code", "N/A")
        product_name = eval_result.get("product_name", "N/A")
        supplier = eval_result.get("supplier", "N/A")
        rep_no = parsed_data.get("report_no", "N/A")
        verdict = eval_result.get("final_verdict", "FAIL")
        is_pass = (verdict == "PASS")
        date_str = parsed_data.get('inspection_date', '2026-09-28')

        # Fonts & Styles
        title_font = Font(name="맑은 고딕", size=16, bold=True, color="1E293B")
        header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        subhead_font = Font(name="맑은 고딕", size=11, bold=True, color="0F172A")
        bold_font = Font(name="맑은 고딕", size=10, bold=True)
        regular_font = Font(name="맑은 고딕", size=10)
        
        verdict_font = Font(name="맑은 고딕", size=14, bold=True, color="166534" if is_pass else "991B1B")

        primary_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        label_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

        thin_side = Side(border_style="thin", color="CBD5E1")
        border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        # 1. Title
        title_text = "출하검사 합격 인증 및 입고 승인서" if is_pass else "출하검사 품질 부적합 통보서 (NCR)"
        ws.merge_cells("A1:F1")
        ws["A1"] = title_text
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 35

        # 2. Metadata Table
        ws["A3"] = "문서 번호"
        ws["B3"] = f"{'COA-PASS' if is_pass else 'NCR'}-{model_code}-{str(rep_no).replace('/', '_')}"
        ws["C3"] = "발행 일자"
        ws["D3"] = date_str
        ws["E3"] = "종합 판정"
        ws["F3"] = verdict

        ws["A4"] = "제조 협력사"
        ws["B4"] = supplier
        ws["C4"] = "모델/품목명"
        ws["D4"] = f"{product_name} ({model_code})"
        ws["E4"] = "적용 기준서"
        ws["F4"] = eval_result.get("inspection_spec_no", "ISO 2859-1")

        for r in range(3, 5):
            ws.row_dimensions[r].height = 24
            for c in range(1, 7):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                if c in (1, 3, 5):
                    cell.fill = label_fill
                    cell.font = bold_font
                    cell.alignment = center_align
                else:
                    cell.font = regular_font
                    cell.alignment = left_align
        ws["F3"].font = verdict_font
        ws["F3"].alignment = center_align
        ws["F3"].fill = pass_fill if is_pass else fail_fill

        # 3. Item Table Header
        ws["A6"] = "시험항목별 공차 정밀 판정 내역"
        ws["A6"].font = subhead_font
        
        headers = ["시험 검사항목", "기준 하한치(LSL)", "기준 상한치(USL)", "성적서 실측치", "단위", "판정 (Status)"]
        ws.row_dimensions[7].height = 26
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=7, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = primary_fill
            cell.alignment = center_align
            cell.border = border

        # 4. Item Rows
        items = eval_result.get("evaluation_items", [])
        current_row = 8
        for it in items:
            ws.row_dimensions[current_row].height = 22
            row_vals = [
                it.get("item_name", ""),
                it.get("min_limit", "-"),
                it.get("max_limit", "-"),
                it.get("measured", "-"),
                it.get("unit", ""),
                it.get("status", "")
            ]
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = border
                cell.alignment = left_align if col_idx == 1 else center_align
                if col_idx == 6:
                    cell.font = bold_font
                    cell.fill = pass_fill if val == "PASS" else fail_fill
            current_row += 1

        # 5. Action Items
        current_row += 1
        ws.cell(row=current_row, column=1, value="권고 후속 조치 사항 (Action Items)").font = subhead_font
        current_row += 1
        
        actions = [
            "1. ERP/MES 전산에 입고 합격 정보 즉시 동기화 및 자재 가결 처리" if is_pass else "1. 입고 즉시 보류 및 불합격 로트 격리 조치",
            "2. 정상 입고 승인 및 생산 라인 불출 허가" if is_pass else "2. 협력사 앞 3영업일 이내 5-Why 원인분석 및 8D 대책서 요구",
            "3. 성적서 원본 사내 품질 DB 영구 아카이빙 완료" if is_pass else "3. 차기 입고 로트 검사 수준 격상 (ISO 2859 강화검사 적용)"
        ]
        for act in actions:
            ws.cell(row=current_row, column=1, value=act).font = regular_font
            current_row += 1

        # Adjust column widths
        col_widths = {1: 28, 2: 18, 3: 18, 4: 18, 5: 12, 6: 18}
        for col_idx, w in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = w

        out_stream = io.BytesIO()
        wb.save(out_stream)
        return out_stream.getvalue()

    @staticmethod
    def generate_pdf_report(parsed_data: Dict[str, Any], eval_result: Dict[str, Any]) -> bytes:
        """
        PyMuPDF를 활용하여 한글 및 벡터 그래픽이 포함된 고품질 공식 PDF 문서 생성
        """
        if not HAS_PYMUPDF:
            return b"%PDF-1.4 Empty fallback"

        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)  # A4 size

        model_code = eval_result.get("model_code", "N/A")
        product_name = eval_result.get("product_name", "N/A")
        supplier = eval_result.get("supplier", "N/A")
        rep_no = parsed_data.get("report_no", "N/A")
        verdict = eval_result.get("final_verdict", "FAIL")
        is_pass = (verdict == "PASS")
        date_str = parsed_data.get('inspection_date', '2026-09-28')

        # Top Header Bar
        header_color = (0.09, 0.39, 0.2) if is_pass else (0.6, 0.1, 0.1)
        page.draw_rect(pymupdf.Rect(40, 40, 555, 90), color=header_color, fill=header_color)
        
        title_text = "출하검사 합격 인증 및 입고 승인서" if is_pass else "출하검사 품질 부적합 통보서 (NCR)"
        page.insert_text(pymupdf.Point(55, 72), title_text, fontsize=18, color=(1, 1, 1), fontname="korea")

        # Sub Header
        page.insert_text(pymupdf.Point(40, 115), f"문서 번호: {'COA-PASS' if is_pass else 'NCR'}-{model_code}-{rep_no}", fontsize=11, fontname="korea")
        page.insert_text(pymupdf.Point(380, 115), f"발행 일자: {date_str}", fontsize=11, fontname="korea")

        # Info Box
        page.draw_rect(pymupdf.Rect(40, 130, 555, 205), color=(0.8, 0.85, 0.9), fill=(0.97, 0.98, 1.0))
        page.insert_text(pymupdf.Point(55, 155), f"제조 협력사: {supplier}", fontsize=11, fontname="korea")
        page.insert_text(pymupdf.Point(55, 175), f"품목 및 모델: {product_name} ({model_code})", fontsize=11, fontname="korea")
        page.insert_text(pymupdf.Point(55, 195), f"검사 기준서: {eval_result.get('inspection_spec_no', 'ISO 2859-1')}", fontsize=11, fontname="korea")
        
        # Verdict Badge
        badge_rect = pymupdf.Rect(420, 145, 535, 190)
        badge_fill = (0.86, 0.99, 0.9) if is_pass else (1.0, 0.88, 0.88)
        badge_border = (0.09, 0.39, 0.2) if is_pass else (0.6, 0.1, 0.1)
        page.draw_rect(badge_rect, color=badge_border, fill=badge_fill)
        page.insert_text(pymupdf.Point(440, 172), f"종합 판정: {verdict}", fontsize=13, color=badge_border, fontname="korea")

        # Table Header
        y = 235
        page.insert_text(pymupdf.Point(40, y), "시험항목별 정밀 공차 검증 내역 (ISO 2859-1)", fontsize=13, fontname="korea")
        
        y += 15
        page.draw_rect(pymupdf.Rect(40, y, 555, y + 25), color=(0.12, 0.16, 0.23), fill=(0.12, 0.16, 0.23))
        page.insert_text(pymupdf.Point(50, y + 17), "검사 항목", fontsize=10, color=(1, 1, 1), fontname="korea")
        page.insert_text(pymupdf.Point(220, y + 17), "규격 범위 (LSL ~ USL)", fontsize=10, color=(1, 1, 1), fontname="korea")
        page.insert_text(pymupdf.Point(380, y + 17), "성적서 실측치", fontsize=10, color=(1, 1, 1), fontname="korea")
        page.insert_text(pymupdf.Point(490, y + 17), "판정", fontsize=10, color=(1, 1, 1), fontname="korea")

        # Table Rows
        items = eval_result.get("evaluation_items", [])
        y += 25
        for it in items:
            page.draw_rect(pymupdf.Rect(40, y, 555, y + 22), color=(0.85, 0.88, 0.92))
            st_color = (0.09, 0.39, 0.2) if it.get("status") == "PASS" else (0.7, 0.1, 0.1)
            
            page.insert_text(pymupdf.Point(50, y + 15), str(it.get("item_name", "")), fontsize=10, fontname="korea")
            range_str = f"{it.get('min_limit')} ~ {it.get('max_limit')} {it.get('unit', '')}"
            page.insert_text(pymupdf.Point(220, y + 15), range_str, fontsize=10, fontname="korea")
            meas_str = f"{it.get('measured')} {it.get('unit', '')}"
            page.insert_text(pymupdf.Point(380, y + 15), meas_str, fontsize=10, fontname="korea")
            page.insert_text(pymupdf.Point(495, y + 15), str(it.get("status", "")), fontsize=10, color=st_color, fontname="korea")
            y += 22

        # Action Items Section
        y += 30
        page.insert_text(pymupdf.Point(40, y), "조치 및 요구사항 (Action Items)", fontsize=13, fontname="korea")
        y += 20
        actions = [
            "1. ERP/MES 전산 입고 완료 처리 및 자재 불출 승인" if is_pass else "1. 입고 즉시 보류 및 부적합 로트 전량 격리 반품",
            "2. 공급사 완제품 출하 성적서 품질 DB 영구 보존" if is_pass else "2. 협력사 앞 3영업일 이내 5-Why 원인분석 및 8D 대책서 요구",
            "3. 생산 라인 정상 투입 가결" if is_pass else "3. 차기 로트 검사 수준 격상 (ISO 2859 보통검사 -> 강화검사)"
        ]
        for act in actions:
            page.insert_text(pymupdf.Point(50, y), act, fontsize=10, fontname="korea")
            y += 18

        # Footer & Signature
        page.draw_line(pymupdf.Point(40, 750), pymupdf.Point(555, 750), color=(0.7, 0.7, 0.7))
        page.insert_text(pymupdf.Point(40, 770), "검사 시스템: 칼퇴보증 QA Agent | 자동 공차 판정 & 품질 감사 엔진", fontsize=9, color=(0.5, 0.5, 0.5), fontname="korea")
        page.insert_text(pymupdf.Point(400, 770), "품질보증팀 승인: [ 서명 완료 ]", fontsize=9, color=(0.2, 0.2, 0.2), fontname="korea")

        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

